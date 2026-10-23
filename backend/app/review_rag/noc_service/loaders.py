"""Bounded, table-aware local loaders. No model calls, formulas or source writes."""
import csv
import io
from zipfile import ZipFile
from xml.etree import ElementTree as ET

MAX_BYTES = 5_000_000
MAX_CHARS = 200_000
MAX_CHUNKS = 100
CHUNK_CHARS = 2800


class Evidence:
    def __init__(self):
        self.chunks = []
        self.characters = 0
        self.tables = 0

    def add(self, locator, text):
        text = text.strip()
        if not text:
            return
        self.characters += len(text)
        if self.characters > MAX_CHARS or len(text) > CHUNK_CHARS:
            raise ValueError('Extracted content exceeds limits; split the document for review')
        self.chunks.append((locator, text))
        if len(locator) > 300 or len(self.chunks) > MAX_CHUNKS:
            raise ValueError('Extraction has too many chunks or an oversized locator')

    def prose(self, label, text):
        lines = text.splitlines()
        buffer, first = [], 1
        for number, line in enumerate(lines, 1):
            if len(line) > CHUNK_CHARS:
                if buffer:self.add(f'{label}, lines {first}-{number-1}', '\n'.join(buffer))
                buffer=[];offset=0
                while offset<len(line):
                    end=min(offset+CHUNK_CHARS,len(line))
                    boundary=line.rfind(' ',offset,end) if end<len(line) else end
                    if boundary>offset+CHUNK_CHARS//2:end=boundary
                    self.add(f'{label}, line {number}, chars {offset+1}-{end}',line[offset:end])
                    offset=end
                first=number+1
                continue
            if buffer and sum(len(part) + 1 for part in buffer) + len(line) > CHUNK_CHARS:
                self.add(f'{label}, lines {first}-{number-1}', '\n'.join(buffer))
                buffer, first = [], number
            buffer.append(line)
        if buffer:
            self.add(f'{label}, lines {first}-{len(lines)}', '\n'.join(buffer))

    def rows(self, label, rows, context=''):
        """Pack complete rows; repeat context across chunks without losing locators."""
        buffer, first, last = [], None, None
        for number, text in rows:
            if not text.strip():
                continue
            if len(context)+len(text)+2>CHUNK_CHARS:
                raise ValueError('A table row exceeds extraction limits; split it for review')
            if buffer and len(context)+sum(len(part)+1 for part in buffer)+len(text)>CHUNK_CHARS:
                self.add(f'{label}, rows {first}-{last}',context+'\n'.join(buffer))
                buffer, first = [], None
            if first is None:first=number
            last=number;buffer.append(text)
        if buffer:self.add(f'{label}, rows {first}-{last}',context+'\n'.join(buffer))

    def table(self, label, rows, *, headers=False):
        """Keep literal row/column positions. Only CSV declares its first row a header."""
        if not rows:
            return
        self.tables+=1
        if max(map(len,rows)) > 100:
            raise ValueError('Table exceeds column limits')
        first=' | '.join(f'C{i+1}={cell_text(value)}' for i,value in enumerate(rows[0]))
        context=f'{label}; '+('headers' if headers else 'first row context (not assumed column headers)')+': '+first+'\n'
        if len(context)>1000:
            # Wide first rows stay in their own cited row rather than multiplying them.
            context=f'{label}; column positions C1, C2, ... follow the source.\n'
        values=[]
        for number,row in enumerate(rows,1):
            if not row:continue
            if headers and len(row)!=len(rows[0]):
                raise ValueError('CSV has inconsistent columns; review extraction')
            if any(value is not None and str(value).strip() for value in row):
                text=f'Row {number}: '+' | '.join(f'C{i+1}={cell_text(value)}' for i,value in enumerate(row))
                if len(context)+len(text)+2>CHUNK_CHARS:
                    self.rows(label,values,context);values=[]
                    for column,value in enumerate(row,1):
                        if value is not None and str(value).strip():self.prose(f'{label}, row {number}, C{column}',cell_text(value))
                else:values.append((number,text))
        self.rows(label,values,context)

    def finish(self):
        if not self.chunks:
            raise ValueError('Document has no extractable text; OCR may be required')
        return self.chunks


def cell_text(value):
    if value is None:
        return '(blank)'
    return value.isoformat() if hasattr(value, 'isoformat') else str(value).strip()


def check_office_zip(raw):
    with ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(item.file_size for item in entries) > 50_000_000:
            raise ValueError('Office archive exceeds extraction limits')
        if any(item.flag_bits & 1 or item.file_size > 20_000_000 for item in entries):
            raise ValueError('Encrypted or oversized Office content requires review')


def pdf_chunks(raw, notes):
    import pdfplumber
    from pypdf import PdfReader
    # Bound decoded page streams before passing text/table detection to pdfminer.
    reader = PdfReader(io.BytesIO(raw), strict=True)
    if reader.is_encrypted or not 1 <= len(reader.pages) <= 100:
        raise ValueError('Encrypted or oversized PDF requires review')
    for page in reader.pages:
        contents = page.get_contents()
        if contents is not None and len(contents.get_data()) > MAX_BYTES:
            raise ValueError('PDF page exceeds extraction limits')
    evidence = Evidence()
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        for number, page in enumerate(pdf.pages, 1):
            if len(page.chars) > MAX_CHARS:
                raise ValueError('PDF page exceeds character limits')
            tables = page.find_tables()
            if page.images:notes.add('pdf_image_content_not_transcribed')
            boxes = [table.bbox for table in tables]
            def outside_tables(obj):
                cx = (obj.get('x0', 0) + obj.get('x1', 0)) / 2
                cy = (obj.get('top', 0) + obj.get('bottom', 0)) / 2
                return not any(x0 <= cx <= x1 and top <= cy <= bottom for x0, top, x1, bottom in boxes)
            text = page.filter(outside_tables).extract_text() or ''
            if not text.strip() and not tables:
                notes.add(f'pdf_page_{number}_requires_ocr_or_blank_review')
            evidence.prose(f'Page {number}', text)
            for index, table in enumerate(tables, 1):
                evidence.table(f'Page {number}, table {index}', table.extract())
    return evidence.finish() if evidence.chunks else []


def docx_chunks(raw, notes):
    from docx import Document
    from docx.table import Table
    check_office_zip(raw)
    document, evidence = Document(io.BytesIO(raw)), Evidence()
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with ZipFile(io.BytesIO(raw)) as archive:
        if any(name.startswith('word/media/') for name in archive.namelist()):notes.add('docx_image_content_not_transcribed')
        if any(name.startswith('word/embeddings/') for name in archive.namelist()):notes.add('docx_embedded_object_not_extracted')
        for name in archive.namelist():
            if not name.startswith('word/') or not name.endswith('.xml'):continue
            if archive.getinfo(name).file_size>8_000_000:raise ValueError('Office XML exceeds extraction limits')
            part=archive.read(name)
            if b'<!DOCTYPE' in part.upper() or b'<!ENTITY' in part.upper():raise ValueError('XML declarations require review')
            element=ET.fromstring(part)
            if element.find('.//w:ins',ns) is not None or element.find('.//w:del',ns) is not None:notes.add('docx_unresolved_tracked_changes')
            if element.find('.//w:txbxContent',ns) is not None:notes.add('docx_textbox_not_extracted')
            if name in ('word/footnotes.xml','word/endnotes.xml'):notes.add('docx_notes_require_review')
    def table_content(table,label,depth=0):
        if depth>3 or len(table.rows)>2000:raise ValueError('Nested or long DOCX table exceeds limits')
        rows=[];nested=[]
        for number,row in enumerate(table.rows,1):
            if len(row.cells)>100:raise ValueError('DOCX table exceeds column limits')
            rows.append([None]*row.grid_cols_before+[cell.text for cell in row.cells]+[None]*row.grid_cols_after)
            seen=set()
            for column,cell in enumerate(row.cells,1):
                if cell._tc in seen:continue
                seen.add(cell._tc)
                for index,child in enumerate(cell.tables,1):nested.append((child,f'{label}, row {number}, C{column}, nested table {index}'))
        evidence.table(label,rows)
        for child,child_label in nested:table_content(child,child_label,depth+1)
    paragraph, table_number = 0, 0
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            table_number += 1
            table_content(block,f'Table {table_number}')
        else:
            paragraph += 1
            evidence.prose(f'Paragraph {paragraph}', block.text)
    # Header/footer text can carry revision information; retain it with a locator.
    seen = set()
    for section_index, section in enumerate(document.sections, 1):
        for kind in ('header', 'footer', 'first_page_header', 'first_page_footer', 'even_page_header', 'even_page_footer'):
            part = getattr(section, kind)
            if part.part.partname in seen:
                continue
            seen.add(part.part.partname)
            evidence.prose(f'Section {section_index} {kind}', '\n'.join(p.text for p in part.paragraphs))
            for index, table in enumerate(part.tables, 1):
                table_content(table,f'Section {section_index} {kind} table {index}')
    return evidence.finish() if evidence.chunks else []


def csv_chunks(raw):
    text = raw.decode('utf-8-sig')
    evidence = Evidence()
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=''), dialect, strict=True)
    rows = []
    for row in reader:
        if len(rows) >= 5000 or len(row) > 100:
            raise ValueError('CSV exceeds row or column limits')
        rows.append(row)
    evidence.table('CSV', rows, headers=True)
    return evidence.finish()


def xlsx_chunks(raw, notes):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    check_office_zip(raw)
    evidence, cells_seen = Evidence(), 0
    formulas = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    cached = None
    try:
        cached = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
        if len(formulas.sheetnames) > 30:
            raise ValueError('Workbook exceeds worksheet limit')
        for sheet in formulas:
            if (sheet.max_row or 0) > 5000 or (sheet.max_column or 0) > 100:
                raise ValueError('Worksheet exceeds row or column limits')
            values = cached[sheet.title]
            row_text=[]
            # Scan beyond potentially incorrect producer-supplied dimensions.
            sheet.reset_dimensions(); values.reset_dimensions()
            for number, (row, cached_row) in enumerate(zip(sheet.iter_rows(), values.iter_rows(), strict=True), 1):
                cells_seen += len(row)
                if number > 5000 or len(row) > 100 or cells_seen > 20000:
                    raise ValueError('Workbook exceeds cell limits')
                entries = []
                for column, (cell, result) in enumerate(zip(row, cached_row, strict=True), 1):
                    if cell.value is None:
                        continue
                    value = cell_text(cell.value)
                    if cell.data_type == 'f':
                        value += ' [cached result: ' + (cell_text(result.value) if result.value is not None else 'unavailable; formula not calculated') + ']'
                        if result.value is None:notes.add('xlsx_formula_cached_result_unavailable')
                    entries.append(f'{get_column_letter(column)}{number}: {value}')
                if entries:
                    row_text.append((number,f'Row {number}: '+' | '.join(entries)))
            context=f'Sheet: {sheet.title} ({sheet.sheet_state})\n'
            if row_text and len(row_text[0][1])<700:context+='First populated row context: '+row_text[0][1]+'\n'
            evidence.rows(f'Sheet {sheet.title!r}',row_text,context)
    finally:
        formulas.close()
        if cached is not None:
            cached.close()
    return evidence.finish()


def assess_document(raw, extension, *, max_bytes=MAX_BYTES):
    extension = extension.lower()
    if not raw or not 1<=max_bytes<=20_000_000 or len(raw) > (200000 if extension == '.txt' else max_bytes):
        raise ValueError('Empty or oversized document')
    notes=set()
    if len(raw)>MAX_BYTES:notes.add('source_exceeds_current_5mb_index_limit')
    if extension == '.pdf':
        if not raw.startswith(b'%PDF-'):
            raise ValueError('Invalid PDF')
        chunks=pdf_chunks(raw,notes)
    elif extension == '.docx':chunks=docx_chunks(raw,notes)
    elif extension == '.csv':chunks=csv_chunks(raw)
    elif extension == '.xlsx':chunks=xlsx_chunks(raw,notes)
    elif extension == '.txt':
        evidence = Evidence()
        evidence.prose('Text', raw.decode('utf-8-sig'))
        chunks=evidence.finish()
        chunks=[(locator.replace('Text, lines','Lines',1),text) for locator,text in chunks]
    else:raise ValueError('Unsupported document format')
    return {'chunks':[{'locator':locator,'text':text} for locator,text in chunks],
        'review_flags':sorted(notes),'status':'needs_review' if notes or not chunks else 'extracted_for_review',
        'approved':False,'indexed':False,'provider_calls':0}


def load_document(raw, extension):
    result=assess_document(raw,extension)
    if result['review_flags'] or not result['chunks']:
        raise ValueError('Source contains unreviewed or unavailable evidence (including image/OCR gaps); inspect extraction flags before indexing')
    return [(chunk['locator'],chunk['text']) for chunk in result['chunks']]
